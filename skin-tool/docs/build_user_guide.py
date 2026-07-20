from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "WorkBuddy初音换肤工具_使用手册_v1.2.docx"
PREVIEW = ROOT / "assets" / "miku" / "miku-reference.jpg"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MIKU_TEAL = "008F8C"
MUTED = "5C6B73"
LIGHT_BLUE = "E8EEF5"
LIGHT_TEAL = "EAF8F7"
LIGHT_GOLD = "FFF7DF"
INK = "13212A"
WHITE = "FFFFFF"


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, size, color=INK, bold=False):
    style.font.name = "Calibri"
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold


def shade_paragraph(paragraph, fill, border=None):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    if border:
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is None:
            p_bdr = OxmlElement("w:pBdr")
            p_pr.append(p_bdr)
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), border)
        p_bdr.append(left)


def add_callout(doc, label, text, fill=LIGHT_TEAL, border=MIKU_TEAL):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.25
    shade_paragraph(paragraph, fill, border)
    label_run = paragraph.add_run(f"  {label}  ")
    set_run_font(label_run, 10.5, True, DARK_BLUE)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, 10.5, False, INK)
    return paragraph


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_label_detail_table(doc, rows, widths_dxa=(1700, 7660)):
    table = doc.add_table(rows=len(rows) + 1, cols=2)
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for index, text in enumerate(("项目", "说明")):
        set_cell_fill(header_cells[index], MIKU_TEAL)
        header_p = header_cells[index].paragraphs[0]
        header_p.paragraph_format.space_after = Pt(0)
        header_run = header_p.add_run(text)
        set_run_font(header_run, 10.5, True, WHITE)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for row_index, (label, detail) in enumerate(rows):
        left, right = table.rows[row_index + 1].cells
        set_cell_fill(left, LIGHT_BLUE)
        left_p = left.paragraphs[0]
        left_p.paragraph_format.space_after = Pt(0)
        left_run = left_p.add_run(label)
        set_run_font(left_run, 10.5, True, DARK_BLUE)
        right_p = right.paragraphs[0]
        right_p.paragraph_format.space_after = Pt(0)
        right_run = right_p.add_run(detail)
        set_run_font(right_run, 10.5, False, INK)
    set_table_geometry(table, list(widths_dxa))
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    set_run_font(run, 11, False, INK)
    return paragraph


def add_step(doc, title, detail, style="List Number"):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.25
    title_run = paragraph.add_run(title)
    set_run_font(title_run, 11, True, INK)
    detail_run = paragraph.add_run(f"：{detail}")
    set_run_font(detail_run, 11, False, INK)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_body(doc, text, bold_prefix=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, 11, True, INK)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest, 11, False, INK)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, 11, False, INK)
    return paragraph


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, 9, False, MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, display, end])
    tail = paragraph.add_run(" 页")
    set_run_font(tail, 9, False, MUTED)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.right_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
set_style_font(normal, 11, INK, False)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

title_style = styles["Title"]
set_style_font(title_style, 30, DARK_BLUE, True)
title_style.paragraph_format.space_before = Pt(0)
title_style.paragraph_format.space_after = Pt(8)

subtitle_style = styles["Subtitle"]
set_style_font(subtitle_style, 14, MUTED, False)
subtitle_style.paragraph_format.space_after = Pt(20)

heading_tokens = {
    1: (16, BLUE, 18, 10),
    2: (13, BLUE, 14, 7),
    3: (12, DARK_BLUE, 10, 5),
}
for level, (size, color, before, after) in heading_tokens.items():
    style = styles[f"Heading {level}"]
    set_style_font(style, size, color, True)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

for name in ("List Bullet", "List Number", "List Number 2"):
    style = styles[name]
    set_style_font(style, 11, INK, False)
    style.paragraph_format.left_indent = Inches(0.375)
    style.paragraph_format.first_line_indent = Inches(-0.188)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.25

header = section.header
header_p = header.paragraphs[0]
header_p.text = "WorkBuddy 换肤工具  |  使用手册"
header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_run_font(header_p.runs[0], 9, True, MUTED)
footer = section.footer
add_page_number(footer.paragraphs[0])

# Cover: editorial_cover pattern with a restrained Miku accent override.
spacer = doc.add_paragraph()
spacer.paragraph_format.space_after = Pt(34)
kicker = doc.add_paragraph()
kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
kicker.paragraph_format.space_after = Pt(14)
set_run_font(kicker.add_run("MIKU MODE 01  ·  QUICK START GUIDE"), 10, True, MIKU_TEAL)
title = doc.add_paragraph("WorkBuddy 初音换肤工具", style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph("安装、深浅模式、恢复与分享使用手册", style="Subtitle")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
if PREVIEW.exists():
    picture_p = doc.add_paragraph()
    picture_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = picture_p.add_run()
    inline = run.add_picture(str(PREVIEW), width=Inches(5.9))
    doc_pr = inline._inline.docPr
    doc_pr.set("descr", "初音未来主题效果参考图")
    picture_p.paragraph_format.space_after = Pt(18)
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.paragraph_format.space_after = Pt(4)
set_run_font(meta.add_run("适用系统：Windows 10/11  |  适配版本：WorkBuddy 5.2.5"), 10.5, True, DARK_BLUE)
note = doc.add_paragraph()
note.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(note.add_run("非官方本地皮肤工具 · 不修改 WorkBuddy 安装包 · 可一键恢复"), 9.5, False, MUTED)

doc.add_page_break()

add_heading(doc, "1. 朋友拿到压缩包后怎么用", 1)
add_callout(doc, "最快用法", "安装 WorkBuddy → 解压整个压缩包 → 双击“打开换肤工具.cmd” → 点击“应用初音皮肤” → 在用户菜单中选择深色或浅色。")
add_step(doc, "确认已安装 WorkBuddy", "先正常打开一次 WorkBuddy，确认能进入主界面后关闭或保持运行。")
add_step(doc, "完整解压工具包", "不要直接在压缩软件预览窗口里运行，也不要只拖出一个 CMD 文件。")
add_step(doc, "打开换肤工具", "双击“打开换肤工具.cmd”。Windows 若弹出安全提示，请先确认压缩包来源确实是可信朋友。")
add_step(doc, "应用初音皮肤", "点击按钮后，工具会提示重启 WorkBuddy；未发送的输入可能丢失，请保存后再继续。")
add_step(doc, "选择外观", "点击 WorkBuddy 左下角头像或用户名，在“外观”中选择“浅色”或“深色”，皮肤会立即跟随切换。")
add_step(doc, "恢复原界面", "再次打开工具，点击“恢复原始界面”，WorkBuddy 会按正常方式重新启动。")

add_heading(doc, "2. 使用前要求", 1)
add_label_detail_table(doc, [
    ("系统", "Windows 10 或 Windows 11，64 位"),
    ("必需软件", "已安装并可正常登录的 WorkBuddy"),
    ("已验证版本", "WorkBuddy 5.2.5；后续版本可能因界面类名变化需要更新适配"),
    ("Node.js", "无需单独安装；系统没有 Node.js 22 时，工具会使用 WorkBuddy 自带的运行时"),
    ("网络", "换肤本身不需要联网；WorkBuddy 自身功能仍按其正常方式联网"),
])

add_heading(doc, "3. 发送给朋友前，请先检查", 1)
add_bullet(doc, "只发送发布版 ZIP，不要把整个开发目录直接打包。")
add_bullet(doc, "不要包含 workbuddy-miku-preview.png；该文件可能显示本机任务名称。")
add_bullet(doc, "告诉对方这是非官方主题工具，来源应为你本人发送的原始 ZIP。")
add_bullet(doc, "公开发布前应替换为拥有明确授权的角色素材；当前素材更适合朋友间本地体验。")

add_heading(doc, "4. 日常操作", 1)
add_heading(doc, "4.1 应用皮肤", 2)
add_body(doc, "打开换肤工具后点击“应用初音皮肤”。若 WorkBuddy 正在运行，工具会询问是否重启。确认后会启动仅监听本机回环地址的调试端口，并运行一个隐藏的注入服务。")
add_bullet(doc, "应用成功后：侧栏、首页、场景按钮和输入框会变成深海青与荧光蓝绿色。")
add_bullet(doc, "真实按钮和输入框仍是 WorkBuddy 原控件，主题装饰层不会拦截鼠标操作。")
add_bullet(doc, "关闭 WorkBuddy 后，注入服务可以继续等待；再次打开工具重新应用即可。")

add_heading(doc, "4.2 切换深色 / 浅色", 2)
add_body(doc, "点击 WorkBuddy 左下角头像或用户名打开用户菜单，在“外观”一行选择“浅色”或“深色”。皮肤会即时切换，不需要重新运行换肤工具；你的选择会在页面跳转、自动化页面和应用重启后保留。")
add_bullet(doc, "浅色模式：薄荷白背景、深色正文，任务标题、菜单、设置和表单均单独适配。")
add_bullet(doc, "深色模式：深海青背景、亮色正文，保留初音主题的荧光青与粉色点缀。")
add_bullet(doc, "恢复原始界面时，皮肤记录的外观偏好会一并清除。")

add_heading(doc, "4.3 检查皮肤", 2)
add_body(doc, "点击“检查皮肤”会检查注入标记、侧栏、首页和输入区，并在工具目录生成一张预览截图。分享工具前，请不要把这张截图一起发送。")

add_heading(doc, "4.4 恢复原始界面", 2)
add_body(doc, "点击“恢复原始界面”后，工具会移除当前页面上的主题、关闭隐藏注入服务、关闭调试会话，并用普通方式重启 WorkBuddy。账号、聊天记录和 WorkBuddy 安装文件不会被删除。")

add_heading(doc, "4.5 手动关闭", 2)
add_body(doc, "如果工具窗口打不开，可先在任务管理器中结束 WorkBuddy，再重新打开“打开换肤工具.cmd”。仍无法恢复时，重新启动电脑后直接启动 WorkBuddy，皮肤不会永久写入安装包。")

add_heading(doc, "5. 文件说明", 1)
add_label_detail_table(doc, [
    ("打开换肤工具.cmd", "朋友双击的入口文件"),
    ("skin-manager.ps1", "图形化管理窗口"),
    ("scripts", "启动、恢复、验证和浏览器注入逻辑"),
    ("assets/miku", "初音主题的 CSS 和参考素材"),
    ("README.md", "简版说明，适合在聊天工具中直接预览"),
    ("WorkBuddy初音换肤工具_使用手册_v1.2.docx", "当前这份完整使用文档"),
], widths_dxa=(3400, 5960))

add_heading(doc, "6. 常见问题", 1)
add_heading(doc, "双击没有反应", 2)
add_bullet(doc, "确认已经完整解压，路径中不要包含尚未解压的 ZIP 虚拟目录。")
add_bullet(doc, "右键“打开换肤工具.cmd”，选择“以普通方式打开”；通常不需要管理员权限。")
add_bullet(doc, "确认 Windows PowerShell 没有被公司安全策略完全禁用。")

add_heading(doc, "提示找不到 WorkBuddy", 2)
add_body(doc, "先从官方渠道安装 WorkBuddy，并正常启动一次。工具默认查找当前用户目录下的 WorkBuddy 安装位置，也会读取常见的 Windows 卸载注册信息。")

add_heading(doc, "应用后还是原界面", 2)
add_bullet(doc, "先点击“恢复原始界面”，再点击“应用初音皮肤”。")
add_bullet(doc, "确认没有多个 WorkBuddy 窗口或升级程序同时运行。")
add_bullet(doc, "WorkBuddy 升级后可能需要新版皮肤适配，请向发送工具的人索取更新版。")

add_heading(doc, "自动化页面又变成深色", 2)
add_body(doc, "请先确认版本信息为 v1.2.0。v1.2 会保存用户在“外观”中的真实选择，并在自动化页面继续保持浅色；若仍异常，先恢复原始界面，再重新应用皮肤。")

add_heading(doc, "为什么会有调试端口", 2)
add_body(doc, "本工具通过 Chromium DevTools Protocol 将 CSS 注入 WorkBuddy 的 Electron 页面。端口只绑定到 127.0.0.1，不对局域网开放；但同一 Windows 用户下的其他本机程序理论上仍可访问，因此不用皮肤时建议点击“恢复原始界面”。")

add_heading(doc, "7. 安全与隐私", 1)
add_callout(doc, "重要", "只运行可信来源的 ZIP。不要接受陌生人发送的二次打包版本，也不要在应用皮肤时同时运行来源不明的软件。", fill=LIGHT_GOLD, border="C28A00")
add_bullet(doc, "工具不修改 WorkBuddy.exe 或 resources/app.asar。")
add_bullet(doc, "工具不会主动读取、上传或打包聊天记录。")
add_bullet(doc, "“检查皮肤”生成的截图可能包含任务名称，仅供本人检查，不要随发布包发送。")
add_bullet(doc, "公开分发角色素材前，应确认相关著作权、商标和二次创作许可。")

add_heading(doc, "8. 卸载工具", 1)
add_step(doc, "先恢复原始界面", "打开工具并点击“恢复原始界面”。", style="List Number 2")
add_step(doc, "删除工具目录", "确认 WorkBuddy 已恢复后，删除解压出来的整个文件夹即可。", style="List Number 2")
add_step(doc, "可选清理缓存", "如需彻底清理，可删除 %LOCALAPPDATA%\\WorkBuddySkin；其中只保存运行状态、日志和自动提取的 Node 运行时。", style="List Number 2")

add_body(doc, "文档版本：1.2  |  生成日期：2026-07-16  |  适配目标：WorkBuddy 5.2.5")

doc.core_properties.title = "WorkBuddy 初音换肤工具使用手册"
doc.core_properties.subject = "安装、应用、恢复、分享与故障排查"
doc.core_properties.creator = ""
doc.core_properties.last_modified_by = ""
doc.core_properties.keywords = "WorkBuddy, 换肤, 初音未来, 使用手册"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
